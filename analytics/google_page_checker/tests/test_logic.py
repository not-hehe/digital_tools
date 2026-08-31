#!/usr/bin/env python3
import asyncio
import sys
import tempfile
import types
from contextlib import asynccontextmanager
from itertools import product
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT))


def _stub_playwright():
    """Подменяет playwright заглушкой, если он не установлен."""
    try:
        import playwright.async_api  # noqa: F401
        return
    except ImportError:
        pass
    pw = types.ModuleType("playwright")
    api = types.ModuleType("playwright.async_api")
    for name in ("Page", "Response", "Browser", "BrowserContext",
                 "Playwright", "Route", "Request", "async_playwright"):
        setattr(api, name, object)
    sys.modules["playwright"] = pw
    sys.modules["playwright.async_api"] = api


_stub_playwright()

import logging  # noqa: E402
logging.disable(logging.CRITICAL)

import re  # noqa: E402
from modules import config as cfg  # noqa: E402
from modules.classifier import classify_site, is_ga_positive  # noqa: E402
from modules.network_listener import (NetworkListener, KIND_COLLECT,  # noqa: E402
                                      KIND_ERROR_TELEMETRY, KIND_OTHER)
import modules.page_worker as page_worker_module  # noqa: E402
page_worker_module.MIN_WAIT_AFTER_LOAD = 0    # в тестах не ждём реальные секунды
page_worker_module.NETWORKIDLE_TIMEOUT = 1    # и затишья сети
page_worker_module.LOAD_STATE_TIMEOUT = 1     # и события load
page_worker_module.TARGETED_WAIT_TIMEOUT = 1  # и прицельного дожидания
page_worker_module.RETRY_PAUSE = 0            # и пауз между попытками
from modules.page_worker import PageWorker, find_html_markers  # noqa: E402
from modules.report_writer import ReportWriter  # noqa: E402


# ---------------------------------------------------------------------------
# Фейковые объекты вместо Playwright
# ---------------------------------------------------------------------------

class FakeRequest:
    def __init__(self, url, method="POST"):
        self.url = url
        self.method = method
        self.post_data = None


class FakeResp:
    def __init__(self, url, status=200, answered=True):
        self.url = url
        self.status = status
        self.answered = answered  # False = ответ на запрос так и не пришёл
        self.request = FakeRequest(url)


class BrokenResp(FakeResp):
    """Ответ, у которого .request падает, — воспроизводит дефект
    playwright-python 1.40 ('dict' object has no attribute '_object')."""

    @property
    def request(self):
        raise AttributeError("'dict' object has no attribute '_object'")

    @request.setter
    def request(self, value):
        self._request = value


class FakeContextRequest:
    """context.request браузера: отвечает на зонд заданным статусом."""

    def __init__(self, page):
        self._page = page

    async def get(self, url, timeout=None):
        self._page.probe_calls.append(url)
        if self._page.probe_error:
            raise RuntimeError(self._page.probe_error)
        return FakeResp(url, self._page.probe_status)


class FakePage:
    """Страница, которая при goto() «отдаёт» заданные ответы и HTML.

    late_responses — ответы, «опаздывающие» к моменту чтения HTML:
    они приходят только во время прицельного дожидания wait_for_event
    (симуляция инжекции скриптов с задержкой, как на платформах с отложенной инжекцией).
    probe_status / probe_error — как отвечает googletagmanager на зонд.
    fail — текст ошибки навигации (True = ERR_NAME_NOT_RESOLVED).
    fail_for — словарь {адрес: текст ошибки}: падать только на этих
    адресах, чтобы проверять лестницу альтернативных адресов.
    """

    def __init__(self, responses, html, fail=False, main_status=200,
                 late_responses=None, probe_status=None, probe_error=None,
                 fail_for=None):
        self._responses = responses
        self._html = html
        self._fail = fail
        self._main_status = main_status
        self._late = list(late_responses or [])
        self._fail_for = dict(fail_for or {})
        self._handlers = {}
        self.goto_calls = 0
        self.goto_urls = []
        self.url = "about:blank"
        self.probe_status = probe_status
        self.probe_error = probe_error
        self.probe_calls = []
        self.context = types.SimpleNamespace(request=FakeContextRequest(self))

    def on(self, event, handler):
        self._handlers[event] = handler

    def remove_listener(self, event, handler):
        self._handlers.pop(event, None)

    async def goto(self, url, **kwargs):
        self.goto_calls += 1
        self.goto_urls.append(url)
        self.url = url
        if self._fail_for:
            error = self._fail_for.get(url)
            if error:
                raise RuntimeError(error)
        elif self._fail:
            raise RuntimeError(self._fail if isinstance(self._fail, str)
                               else "net::ERR_NAME_NOT_RESOLVED")
        req_handler = self._handlers.get("request")
        resp_handler = self._handlers.get("response")
        for resp in self._responses:
            if req_handler:
                await req_handler(resp.request)     # запрос отправлен...
            if resp_handler and resp.answered:
                await resp_handler(resp)            # ...ответ пришёл (или нет)
        return FakeResp(url, self._main_status)     # ответ главной страницы

    async def content(self):
        return self._html

    async def wait_for_load_state(self, state="load", timeout=None):
        pass  # у фейковой страницы сеть «затихает» мгновенно

    async def wait_for_event(self, event, predicate=None, timeout=None):
        """Прицельное дожидание: «доставляет» опоздавшие ответы через
        хендлеры listener'а и возвращает первый подошедший под predicate.
        Если ничего не подошло — таймаут, как в реальном Playwright."""
        req_handler = self._handlers.get("request")
        resp_handler = self._handlers.get("response")
        matched = None
        for resp in self._late:
            if req_handler:
                await req_handler(resp.request)
            if resp_handler and resp.answered:
                await resp_handler(resp)
            if (matched is None and event == "response" and resp.answered
                    and (predicate is None or predicate(resp))):
                matched = resp
        self._late = []
        if matched is not None:
            return matched
        raise TimeoutError(f"событие {event} не наступило")


class FakeBM:
    def __init__(self, page):
        self._page = page
        self.contexts_opened = 0

    @asynccontextmanager
    async def page_context(self):
        self.contexts_opened += 1
        yield self._page


# ---------------------------------------------------------------------------
# Общие образцы данных
# ---------------------------------------------------------------------------

GTM_HTML = """<html><head>
<script>(function(w,d,s,l,i){...})(window,document,'script','dataLayer','GTM-AB12CD');</script>
<script async src="https://www.googletagmanager.com/gtag/js?id=G-XYZ98765"></script>
<script>gtag('config', 'G-XYZ98765');</script>
</head><body>
<noscript><iframe src="https://www.googletagmanager.com/ns.html?id=GTM-AB12CD"></iframe></noscript>
</body></html>"""

CLEAN_HTML = "<html><body><p>Страница без аналитики</p></body></html>"

# Только тег Universal Analytics — самый частый случай в боевом списке:
# наследие, оставшееся в шаблоне платформы после отключения UA.
UA_HTML = ("<html><head><script async "
           "src=\"https://www.googletagmanager.com/gtag/js?id=UA-10659796-44\">"
           "</script><script>gtag('config','UA-10659796-44');</script>"
           "</head><body>x</body></html>")

# И GTM, и UA в одной вёрстке: живая установка должна победить наследие.
GTM_AND_UA_HTML = ("<html><head>"
                   "<script>(function(){})('GTM-AB12CD');</script>"
                   "<script>gtag('config','UA-10659796-44');</script>"
                   "</head><body>x</body></html>")

# Слабое упоминание: чужой скрипт проверяет наличие gtag, ID нигде нет
WEAK_HTML = ("<html><body><script>"
             "if (window.gtag) { gtag('event', 'from_widget'); }"
             "</script></body></html>")

COLLECT = FakeResp("https://region1.google-analytics.com/g/collect?v=2&tid=G-XYZ98765", 204)
ERRTEL = FakeResp("https://www.google-analytics.com/u/d?t=error&_e=exc&_v=j101&_m=TypeError", 200)
GTM404 = FakeResp("https://www.googletagmanager.com/gtm.js?id=GTM-DEAD01", 404)
NOT_GOOGLE = FakeResp("https://example.com/style.css", 200)
# Хит, отправленный в никуда: endpoint не ответил (блокировка/обрыв сети)
COLLECT_NOANSWER = FakeResp("https://www.google-analytics.com/g/collect?v=2&tid=G-SILENT1",
                            answered=False)


# ---------------------------------------------------------------------------
# Тесты
# ---------------------------------------------------------------------------

def test_config_patterns():
    """Regex-группы конфига: компиляция и разделение хиты/телеметрия/прочее."""
    for group in (cfg.GA_PATTERNS, cfg.GA_COLLECT_PATTERNS, cfg.HTML_GTM_PATTERNS):
        for p in group:
            re.compile(p)

    real_hits = [
        "https://www.google-analytics.com/collect?v=1&t=pageview",
        "https://www.google-analytics.com/g/collect?v=2",
        "https://region1.google-analytics.com/g/collect?v=2",
        "https://www.google-analytics.com/j/collect?v=1",
        "https://www.google-analytics.com/mp/collect?api_secret=x",
        "https://www.google-analytics.com/batch",
        "https://analytics.google.com/g/collect?v=2",
    ]
    for url in real_hits:
        assert any(re.search(p, url) for p in cfg.GA_COLLECT_PATTERNS), url

    not_hits = [
        "https://www.google-analytics.com/u/d?t=error&_e=exc",
        "https://www.googletagmanager.com/gtm.js?id=GTM-AB12CD",
        "https://www.googletagmanager.com/gtag/js?id=G-XYZ98765",
    ]
    for url in not_hits:
        assert not any(re.search(p, url) for p in cfg.GA_COLLECT_PATTERNS), url

    assert all(m in ERRTEL.url for m in cfg.GA_ERROR_MARKERS)

    # секрета в конфиге быть не должно: вебхук приходит только из окружения
    source = (PROJECT_ROOT / "modules" / "config.py").read_text(encoding="utf-8")
    assert "/hooks/" not in source, "в config.py остался адрес вебхука"
    assert "GPC_MATTERMOST_WEBHOOK" in source


def test_classifier():
    """Полный перебор всех 8 комбинаций признаков."""
    expected = {
        (True, True, True): cfg.STATUS_WORKING,
        (True, True, False): cfg.STATUS_WORKING,
        (True, False, True): cfg.STATUS_WORKING,
        (True, False, False): cfg.STATUS_WORKING,
        (False, True, True): cfg.STATUS_BROKEN,
        (False, True, False): cfg.STATUS_BROKEN,
        (False, False, True): cfg.STATUS_SUSPICIOUS,
        (False, False, False): cfg.STATUS_NONE,
    }
    for combo in product([True, False], repeat=3):
        assert classify_site(*combo) == expected[combo], combo

    assert is_ga_positive(cfg.STATUS_WORKING)
    assert is_ga_positive(cfg.STATUS_BROKEN)
    assert not is_ga_positive(cfg.STATUS_SUSPICIOUS)
    assert not is_ga_positive(cfg.STATUS_NONE)
    assert not is_ga_positive(cfg.STATUS_ERROR)


def test_failures_taxonomy():
    """Разбор текста ошибки на вид отказа + политика повторов."""
    from modules import failures as f

    cases = {
        "net::ERR_NAME_NOT_RESOLVED at https://x.ru/": f.FAIL_DNS,
        "net::ERR_SSL_VERSION_OR_CIPHER_MISMATCH at https://x.ru/": f.FAIL_TLS,
        "net::ERR_SSL_UNRECOGNIZED_NAME_ALERT": f.FAIL_TLS,
        "net::ERR_CONNECTION_TIMED_OUT at https://x.ru/": f.FAIL_TIMEOUT,
        "Timeout 45000ms exceeded.": f.FAIL_TIMEOUT,
        "net::ERR_CONNECTION_REFUSED": f.FAIL_REFUSED,
        "net::ERR_CONNECTION_RESET": f.FAIL_REFUSED,
        "net::ERR_EMPTY_RESPONSE": f.FAIL_REFUSED,
        # прерванная навигация — не «порт закрыт»: перебирать адреса
        # и советовать «проверьте порт» здесь неправильно
        "net::ERR_ABORTED": f.FAIL_BROWSER,
        "net::ERR_HTTP_RESPONSE_CODE_FAILURE": f.FAIL_HTTP,
        "net::ERR_TOO_MANY_REDIRECTS": f.FAIL_REDIRECT_LOOP,
        "net::ERR_INVALID_AUTH_CREDENTIALS": f.FAIL_AUTH_REQUIRED,
        "net::ERR_TUNNEL_CONNECTION_FAILED": f.FAIL_PROXY,
        "net::ERR_PROXY_CONNECTION_FAILED": f.FAIL_PROXY,
        "Protocol error (Page.navigate): Cannot navigate to invalid URL":
            f.FAIL_INVALID_URL,
        "'dict' object has no attribute '_object'": f.FAIL_BROWSER,
        "что-то невиданное": f.FAIL_OTHER,
    }
    for text, kind in cases.items():
        assert f.classify_failure(text) == kind, text
    assert f.classify_failure("") == f.FAIL_OTHER

    # у каждого вида есть человекочитаемая метка и подсказка
    for kind in f.FAILURE_ORDER:
        assert f.failure_label(kind) and f.failure_hint(kind)

    # повторять имеет смысл только транзиентные отказы: три попытки
    # по 45 секунд на мёртвую DNS-запись и съедали основное время прогона
    assert f.is_retryable(f.FAIL_TIMEOUT)
    assert f.is_retryable(f.FAIL_REFUSED)
    assert f.is_retryable(f.FAIL_BROWSER)
    assert not f.is_retryable(f.FAIL_DNS)
    assert not f.is_retryable(f.FAIL_INVALID_URL)
    assert not f.is_retryable(f.FAIL_AUTH_REQUIRED)

    # менять адрес имеет смысл там, где протокол или имя хоста могут спасти
    assert f.worth_alternate_url(f.FAIL_TLS)
    assert f.worth_alternate_url(f.FAIL_DNS)
    # молча дропнутый на файрволе 443 выглядит как таймаут, а 80 открыт
    assert f.worth_alternate_url(f.FAIL_TIMEOUT)
    assert not f.worth_alternate_url(f.FAIL_INVALID_URL)
    assert not f.worth_alternate_url(f.FAIL_AUTH_REQUIRED)
    assert not f.worth_alternate_url(f.FAIL_BROWSER)

    # каждый вид, объявленный «стоит сменить адрес», должен реально
    # порождать кандидатов — иначе ветка в url_variants мертва
    from modules.url_variants import alternate_urls
    for kind in f.WORTH_ALTERNATE_URL:
        assert alternate_urls("https://site.ru", kind), kind


def test_url_variants():
    """Нормализация адреса, отсев мусора и подбор альтернатив."""
    from modules import failures as f
    from modules.url_variants import normalize, looks_like_url, alternate_urls

    assert normalize("example.com") == "https://example.com"
    assert normalize("http://example.com") == "http://example.com"
    assert normalize("  example.com  ") == "https://example.com"

    # мусорные ячейки Excel до браузера доходить не должны
    for junk in ("нет", "lkb2b", "уточняется", "-", ""):
        assert not looks_like_url(junk), junk
    for good in ("example.com", "https://a.b.example.com/path?x=1", "localhost:8080",
                 "10.0.0.1", "http://xn--80a.xn--p1ai"):
        assert looks_like_url(good), good

    # TLS не согласовался -> пробуем http, потом другое имя хоста
    alts = alternate_urls("https://old.example.com", f.FAIL_TLS)
    assert alts[0] == "http://old.example.com"
    assert "https://www.old.example.com" in alts

    # DNS от протокола не зависит: меняем только имя
    assert alternate_urls("https://example.com", f.FAIL_DNS) == ["https://www.example.com"]
    assert alternate_urls("https://www.example.com", f.FAIL_DNS) == ["https://example.com"]

    # путь и параметры сохраняются
    alts = alternate_urls("https://example.com/auth?x=1", f.FAIL_REFUSED)
    assert alts[0] == "http://example.com/auth?x=1"

    # мусор и требование авторизации альтернатив не порождают
    assert alternate_urls("https://example.com", f.FAIL_INVALID_URL) == []
    assert alternate_urls("https://example.com", f.FAIL_AUTH_REQUIRED) == []

    # www перед IP-адресом или логином не имеет смысла и ломает адрес:
    # такой кандидат не должен появляться вовсе
    assert alternate_urls("https://10.0.0.5", f.FAIL_DNS) == []
    assert alternate_urls("https://10.0.0.5/lk", f.FAIL_REFUSED) == \
        ["http://10.0.0.5/lk"]
    for alt in alternate_urls("https://user:p@site.ru", f.FAIL_REFUSED):
        assert "www.user" not in alt, alt


def test_network_listener():
    """Разметка запросов по kind, дозаполнение статусов, запрос без ответа."""
    listener = NetworkListener(page=None, ga_patterns=cfg.GA_PATTERNS)
    loop = asyncio.new_event_loop()
    for resp in (COLLECT, ERRTEL, GTM404, NOT_GOOGLE):
        loop.run_until_complete(listener._handle_request(resp.request))
        loop.run_until_complete(listener._handle_response(resp))
    # хит, на который ответ так и не пришёл: только событие request
    loop.run_until_complete(listener._handle_request(COLLECT_NOANSWER.request))
    loop.close()

    assert len(listener.get_all_requests()) == 4          # css отсеян
    assert len(listener.get_collect_requests()) == 2
    assert len(listener.get_error_requests()) == 1
    assert len(listener.get_ga_requests()) == 3           # всё, кроме телеметрии
    kinds = {r["kind"] for r in listener.get_all_requests()}
    assert kinds == {KIND_COLLECT, KIND_ERROR_TELEMETRY, KIND_OTHER}

    dead = [r for r in listener.get_ga_requests() if r["kind"] == KIND_OTHER]
    assert dead[0]["status"] == 404                        # мёртвый gtm.js виден
    answered = [r for r in listener.get_collect_requests() if r["status"] == 204]
    silent = [r for r in listener.get_collect_requests() if r["status"] is None]
    assert len(answered) == 1 and len(silent) == 1         # попытка без ответа учтена


def test_network_listener_broken_response():
    """Дефект playwright 1.40: response.request падает.

    Статус всё равно должен проставиться — по URL ответа, а исключение
    не должно уходить наружу и ронять попытку загрузки сайта.
    """
    listener = NetworkListener(page=None, ga_patterns=cfg.GA_PATTERNS)
    url = "https://www.googletagmanager.com/gtm.js?id=GTM-AB12CD"
    broken = BrokenResp(url, 404)
    loop = asyncio.new_event_loop()
    loop.run_until_complete(listener._handle_request(FakeRequest(url)))
    loop.run_until_complete(listener._handle_response(broken))  # не падает
    loop.close()

    requests = listener.get_ga_requests()
    assert len(requests) == 1
    assert requests[0]["status"] == 404, "статус не восстановлен по URL"


def test_find_html_markers():
    """Поиск маркеров в вёрстке: находки, сниппеты, дедуп, чистый HTML."""
    markers = find_html_markers(GTM_HTML)
    values = {m["marker"] for m in markers}
    assert "GTM-AB12CD" in values
    assert any("G-XYZ98765" in v for v in values)
    assert all(m["snippet"] for m in markers)
    strengths = {m["marker"]: m["strength"] for m in markers}
    assert strengths["GTM-AB12CD"] == "strong"    # ID установки — сильный
    assert strengths["gtag("] == "weak"           # голый вызов — слабый

    weak_only = find_html_markers(WEAK_HTML)
    assert weak_only and all(m["strength"] == "weak" for m in weak_only)

    assert find_html_markers(CLEAN_HTML) == []
    assert find_html_markers("") == []
    assert len(find_html_markers("GTM-SAME01 " * 10)) == 1  # дедупликация

    # лимит на паттерн считает РАЗНЫЕ значения: повторы одного ID не должны
    # вытеснять следующие контейнеры из отчёта
    many = "GTM-AAAA01 " * 5 + "GTM-BBBB02 GTM-CCCC03"
    found = {m["marker"] for m in find_html_markers(many)}
    assert found == {"GTM-AAAA01", "GTM-BBBB02", "GTM-CCCC03"}, found


def _run_worker(page, url="example.com"):
    return asyncio.run(PageWorker(FakeBM(page)).check_pageview_only(url))


def _real_gotos(page):
    """Навигации без служебных сбросов на about:blank."""
    return [u for u in page.goto_urls if u != "about:blank"]


def test_page_worker():
    """Основные сценарии: WORKING/BROKEN(x2)/SUSPICIOUS/NONE/ERROR."""
    r = _run_worker(FakePage([COLLECT, GTM404], GTM_HTML))
    assert r["status"] == cfg.STATUS_WORKING and r["has_ga"]

    r = _run_worker(FakePage([GTM404, ERRTEL], GTM_HTML))
    assert r["status"] == cfg.STATUS_BROKEN and r["has_ga"]

    r = _run_worker(FakePage([], GTM_HTML))       # тег есть, в сети тишина
    assert r["status"] == cfg.STATUS_BROKEN and r["has_ga"]
    assert r["container_id"] == "GTM-AB12CD"

    r = _run_worker(FakePage([ERRTEL], CLEAN_HTML))
    assert r["status"] == cfg.STATUS_SUSPICIOUS and not r["has_ga"]

    # хит отправлен, но endpoint не ответил — аналитика на сайте есть
    r = _run_worker(FakePage([COLLECT_NOANSWER], CLEAN_HTML))
    assert r["status"] == cfg.STATUS_WORKING and r["has_ga"]

    r = _run_worker(FakePage([], CLEAN_HTML))
    assert r["status"] == cfg.STATUS_NONE and not r["has_ga"]

    # только слабые упоминания в вёрстке, сети нет -> NONE, а не BROKEN
    r = _run_worker(FakePage([], WEAK_HTML))
    assert r["status"] == cfg.STATUS_NONE and not r["has_ga"]
    assert any(m["strength"] == "weak" for m in r["html_markers"])

    r = _run_worker(FakePage([], CLEAN_HTML, fail="net::ERR_CONNECTION_RESET"))
    assert r["status"] == cfg.STATUS_ERROR
    assert r["url"] == "https://example.com"       # протокол дополнен
    assert r["failure_kind"] == "refused"


def test_no_retry_on_permanent_failure():
    """Детерминированные отказы не должны отрабатывать три попытки.

    Раньше MAX_RETRIES применялся ко всему подряд: 84 мёртвых по DNS
    адреса × 3 попытки × пауза 5 секунд — и это в каждом из прогонов.
    """
    page = FakePage([], CLEAN_HTML, fail="net::ERR_NAME_NOT_RESOLVED")
    r = _run_worker(page)
    assert r["status"] == cfg.STATUS_ERROR
    assert r["failure_kind"] == "dns"
    # одна попытка на исходный адрес + одна на www-вариант, без повторов
    assert _real_gotos(page) == ["https://example.com",
                                 "https://www.example.com"], page.goto_urls

    # транзиентный отказ по-прежнему добивается повторами
    page = FakePage([], CLEAN_HTML, fail="net::ERR_CONNECTION_TIMED_OUT")
    _run_worker(page)
    assert page.goto_urls.count("https://example.com") == cfg.MAX_RETRIES

    # мусорная строка вообще не доходит до браузера
    bm = FakeBM(FakePage([], CLEAN_HTML))
    r = asyncio.run(PageWorker(bm).check_pageview_only("нет"))
    assert r["status"] == cfg.STATUS_ERROR
    assert r["failure_kind"] == "invalid_url"
    assert bm.contexts_opened == 0, "браузер открывали ради мусорной строки"


def test_alternate_url_ladder():
    """Сайт, недоступный по https, открывается по http — и это не ошибка."""
    page = FakePage([], GTM_HTML, fail_for={
        "https://old.example.com": "net::ERR_SSL_VERSION_OR_CIPHER_MISMATCH",
    })
    r = _run_worker(page, "old.example.com")
    assert r["status"] == cfg.STATUS_BROKEN, r["status"]
    assert r["opened_url"] == "http://old.example.com"
    assert r["url"] == "https://old.example.com"          # идентичность сохранена
    assert r["tried_urls"] == ["https://old.example.com", "http://old.example.com"]

    # DNS есть только у www-варианта
    page = FakePage([], CLEAN_HTML, fail_for={
        "https://shop.ru": "net::ERR_NAME_NOT_RESOLVED",
    })
    r = _run_worker(page, "shop.ru")
    assert r["status"] == cfg.STATUS_NONE
    assert r["opened_url"] == "https://www.shop.ru"

    # число вариантов ограничено: бесконечно перебирать адреса нельзя
    page = FakePage([], CLEAN_HTML, fail="net::ERR_CONNECTION_REFUSED")
    r = _run_worker(page, "dead.ru")
    assert r["status"] == cfg.STATUS_ERROR
    assert len(r["tried_urls"]) <= page_worker_module._MAX_URL_VARIANTS


def test_page_reset_between_variants():
    """Страница гасится перед КАЖДЫМ следующим адресом.

    Chromium после неудачной навигации ещё коммитит свою chrome-error-
    страницу и прерывает ею следующий goto. Сброс должен стоять на входе
    в итерацию, а не в ветке постановки альтернатив: до второго варианта
    можно дойти после отказа, который сам альтернатив не порождает
    (здесь — требование авторизации).
    """
    page = FakePage([], CLEAN_HTML, fail_for={
        "https://site.ru": "net::ERR_CONNECTION_REFUSED",
        "http://site.ru": "net::ERR_INVALID_AUTH_CREDENTIALS",
    })
    r = _run_worker(page, "site.ru")
    assert r["status"] == cfg.STATUS_NONE
    assert page.goto_urls == (
        # refused транзиентен — добивается повторами
        ["https://site.ru"] * cfg.MAX_RETRIES
        + ["about:blank", "http://site.ru"]   # авторизация: одна попытка
        + ["about:blank", "https://www.site.ru"]
    ), page.goto_urls

    # после ПОСЛЕДНЕГО варианта сбрасывать нечего — лишние секунды
    # на каждый безнадёжный сайт в каждом прогоне
    page = FakePage([], CLEAN_HTML, fail="net::ERR_NAME_NOT_RESOLVED")
    _run_worker(page, "dead.ru")
    assert page.goto_urls[-1] != "about:blank", page.goto_urls


def test_http_error_pages_are_analyzed():
    """Сервер ответил кодом ошибки — страница есть, разбираем содержимое."""
    # 404 стабилен: повторять бессмысленно, разбираем сразу
    page = FakePage([], GTM_HTML, main_status=404)
    r = _run_worker(page)
    assert r["status"] == cfg.STATUS_BROKEN, r["status"]
    assert r["http_error"] and r["http_status"] == 404
    assert _real_gotos(page) == ["https://example.com"], \
        "стабильный 4xx не должен ретраиться"

    # 403 — то же самое
    r = _run_worker(FakePage([COLLECT], GTM_HTML, main_status=403))
    assert r["status"] == cfg.STATUS_WORKING and r["http_error"]

    # 503 транзиентен: добиваем повторами, но результат всё равно разбираем
    page = FakePage([], GTM_HTML, main_status=503)
    r = _run_worker(page)
    assert page.goto_calls == cfg.MAX_RETRIES
    assert r["status"] == cfg.STATUS_BROKEN and r["http_status"] == 503

    # обычный ответ пометки не получает
    r = _run_worker(FakePage([], GTM_HTML))
    assert not r["http_error"]


def test_report_writer():
    """Отчёт: все секции, маркеры в BROKEN, совместимость со старым форматом."""
    results = [
        _run_worker(FakePage([COLLECT], GTM_HTML)),
        _run_worker(FakePage([GTM404, ERRTEL], GTM_HTML)),
        _run_worker(FakePage([ERRTEL], CLEAN_HTML)),
        _run_worker(FakePage([], CLEAN_HTML)),
        _run_worker(FakePage([], WEAK_HTML)),
        # gudok при молчащей сети: улику добыл зонд
        _run_worker(FakePage([], GTM_HTML, probe_status=404)),
        _run_worker(FakePage([], CLEAN_HTML, fail=True)),
        # словарь старого формата (без 'status') — из аварийного фолбэка
        {"url": "https://legacy.ru", "has_ga": False, "ga_requests": [],
         "error": "RuntimeError: boom"},
    ]
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / "report.txt"
        ReportWriter(out).write_results_simple(results, run_count=3)
        text = out.read_text(encoding="utf-8")

    for fragment in ("Циклов проверки              : 3",
                     "--- WORKING", "--- BROKEN", "--- SUSPICIOUS",
                     "--- NONE", "--- ERROR", "КОНЕЦ ОТЧЁТА",
                     "GTM-AB12CD", "контекст:", "[HTTP 404]",
                     "Статус : BROKEN — контейнер удалён (gtm.js → 404)",
                     "Google : Found (сетевой запрос со страницы — "
                     "в мини-отчёте)",
                     "Зонд   : https://www.googletagmanager.com/gtm.js"
                     "?id=GTM-AB12CD → HTTP 404",
                     "слабые упоминания в вёрстке",
                     "со слабыми упоминаниями : 1",
                     "https://legacy.ru",
                     # ошибки разложены по виду отказа, а не одной кучей
                     "[dns]", "[other]", "что делать:"):
        assert fragment in text, f"в отчёте нет: {fragment}"
    assert "Google Analytics есть        : 3" in text
    # подсветку получают ровно сайты с сетевой уликой (критерий мини-отчёта):
    # WORKING (collect) и BROKEN с запросом gtm.js; probe-BROKEN — нет
    assert text.count("Google : Found") == 2

    # без run_count строка циклов не выводится (обратная совместимость)
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / "report.txt"
        ReportWriter(out).write_results_simple(results)
        assert "Циклов проверки" not in out.read_text(encoding="utf-8")


def test_container_grouping():
    """Один контейнер на несколько сайтов — отдельной сводкой.

    В боевом прогоне 24 сайта делили GTM-W9D3BC и лежали в отчёте
    двадцатью четырьмя независимыми записями.
    """
    from modules.report_writer import group_by_container

    results = [_run_worker(FakePage([], GTM_HTML), f"site{i}.ru")
               for i in range(3)]
    results.append(_run_worker(FakePage([], CLEAN_HTML), "clean.ru"))

    grouped = group_by_container(results)
    assert list(grouped) == ["GTM-AB12CD"]
    assert len(grouped["GTM-AB12CD"]) == 3

    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / "report.txt"
        ReportWriter(out).write_results_simple(results)
        text = out.read_text(encoding="utf-8")
    assert "Общие контейнеры" in text
    assert "GTM-AB12CD — сайтов: 3" in text
    for i in range(3):
        assert f"https://site{i}.ru" in text


def test_universal_analytics_is_dead():
    """UA-тег — мёртвый код, а не «выключенная конфигом» аналитика.

    Google остановил сбор данных в Universal Analytics 01.07.2023 и удалил
    ресурсы 01.07.2024. Зонд по UA-идентификатору при этом возвращает 200,
    потому что googletagmanager отдаёт gtag.js на любой валидный по форме ID.
    Раньше из этого делался вывод «контейнер жив, страница его не
    запрашивает» — на боевом прогоне неверная формулировка для 90 сайтов
    из 120.
    """
    from modules.classifier import is_universal_analytics
    from modules.page_worker import (detect_broken_reason, extract_container_id,
                                     _extract_probe_target)

    assert is_universal_analytics("UA-10659796-44")
    assert not is_universal_analytics("GTM-AB12CD")
    assert not is_universal_analytics("G-XYZ98765")
    assert not is_universal_analytics(None)

    ua = [{"marker": "UA-10659796-44", "strength": "strong"}]
    gtm_and_ua = ua + [{"marker": "GTM-AB12CD", "strength": "strong"}]

    # живая установка приоритетнее наследия
    assert extract_container_id(ua) == "UA-10659796-44"
    assert extract_container_id(gtm_and_ua) == "GTM-AB12CD"

    # зондировать UA бессмысленно — цели нет, запрос не уходит
    assert _extract_probe_target(ua) is None
    assert _extract_probe_target(gtm_and_ua)[0] == "GTM-AB12CD"

    # вердикт по UA не зависит от того, что ответила сеть
    assert detect_broken_reason([], None, "UA-10659796-44") == cfg.BROKEN_UA_SUNSET
    assert detect_broken_reason([], {"status": 200},
                                "UA-10659796-44") == cfg.BROKEN_UA_SUNSET
    # но собственный запрос страницы к контейнеру сильнее: если страница
    # реально сходила за gtm.js и получила 404, это точнее
    gtm404 = [{"url": "https://www.googletagmanager.com/gtm.js?id=GTM-X",
               "status": 404}]
    assert detect_broken_reason(gtm404, None,
                                "UA-10659796-44") == cfg.BROKEN_CONTAINER_DELETED

    # интеграция: сайт только с UA-тегом
    page = FakePage([], UA_HTML, probe_status=200)
    r = _run_worker(page)
    assert r["status"] == cfg.STATUS_BROKEN
    assert r["container_id"] == "UA-10659796-44"
    assert r["broken_reason"] == cfg.BROKEN_UA_SUNSET
    assert not page.probe_calls, "по UA зонд ходить не должен"

    # сайт с GTM и UA сразу — разбирается как GTM, зонд уходит
    page = FakePage([], GTM_AND_UA_HTML, probe_status=200)
    r = _run_worker(page)
    assert r["container_id"] == "GTM-AB12CD"
    assert r["broken_reason"] == cfg.BROKEN_ALIVE_NOT_LOADED
    assert page.probe_calls


def test_ua_in_reports():
    """Мёртвый UA выделен в отчётах отдельной секцией и счётчиком."""
    from modules.report_writer import count_ua_sunset

    results = [_run_worker(FakePage([], UA_HTML), f"ua{i}.ru") for i in range(3)]
    results.append(_run_worker(FakePage([], GTM_HTML, probe_status=200), "gtm.ru"))
    results.append(_run_worker(FakePage([], CLEAN_HTML), "clean.ru"))

    assert count_ua_sunset(results) == 3

    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / "report.txt"
        writer = ReportWriter(out)
        writer.write_results_simple(results, run_count=1)
        text = out.read_text(encoding="utf-8")
        summary = writer.build_summary(results)

    assert "из них мёртвый тег Universal Analytics : 3" in text
    assert "--- BROKEN — мёртвый тег Universal Analytics" in text
    assert "Google отключил систему, чинить нечего" in text
    assert "устаревший тег Universal Analytics" in text
    # общий UA-контейнер помечен прямо в сводке по контейнерам
    assert "UA-10659796-44 — сайтов: 3" in text
    assert "Universal Analytics, система отключена Google" in text
    # живой GTM в свою секцию не съехал
    assert "--- BROKEN — GA в вёрстке есть, рабочих хитов нет (1) ---" in text

    assert "мёртвый тег Universal Analytics (отключён Google) : 3" in summary


def test_mattermost_webhook_source():
    """Вебхук берётся из окружения, иначе из secrets/ — но не из кода."""
    import importlib
    import os
    from modules import config as live_cfg

    source = (PROJECT_ROOT / "modules" / "config.py").read_text(encoding="utf-8")
    assert "/hooks/" not in source, "адрес вебхука не должен лежать в коде"

    gitignore = (PROJECT_ROOT / ".gitignore").read_text(encoding="utf-8")
    assert "secrets/" in gitignore, "каталог с секретами не закрыт от git"

    original = os.environ.get("GPC_MATTERMOST_WEBHOOK")
    try:
        os.environ["GPC_MATTERMOST_WEBHOOK"] = "https://env.example/hooks/abc"
        assert importlib.reload(live_cfg).MATTERMOST_WEBHOOK_URL == \
            "https://env.example/hooks/abc", "окружение должно иметь приоритет"
        os.environ.pop("GPC_MATTERMOST_WEBHOOK")
        from_file = importlib.reload(live_cfg).MATTERMOST_WEBHOOK_URL
        if live_cfg.MATTERMOST_WEBHOOK_FILE.exists():
            assert from_file.startswith("http"), "файл секрета не прочитался"
            assert not from_file.startswith("#"), "комментарий принят за адрес"
    finally:
        if original is None:
            os.environ.pop("GPC_MATTERMOST_WEBHOOK", None)
        else:
            os.environ["GPC_MATTERMOST_WEBHOOK"] = original
        importlib.reload(live_cfg)


def test_broken_reason():
    """Детализация BROKEN по HTTP-статусу загрузки контейнера."""
    from modules.page_worker import detect_broken_reason

    def req(url, status):
        return {"url": url, "status": status, "kind": "other_google"}

    gtm = "https://www.googletagmanager.com/gtm.js?id=GTM-AB12CD"
    assert detect_broken_reason([req(gtm, 404)]) == cfg.BROKEN_CONTAINER_DELETED
    assert detect_broken_reason([req(gtm, 200)]) == cfg.BROKEN_LOADS_NO_HITS
    assert detect_broken_reason([req(gtm, None)]) == cfg.BROKEN_NO_RESPONSE
    assert detect_broken_reason([]) == cfg.BROKEN_NOT_LOADED
    # 404 сильнее 200: хоть один мёртвый лоадер — контейнер удалён
    assert detect_broken_reason(
        [req(gtm, 200), req(gtm, 404)]) == cfg.BROKEN_CONTAINER_DELETED

    # интеграция: страница как shop.example.com (тег в вёрстке, gtm.js -> 404)
    r = _run_worker(FakePage([GTM404], GTM_HTML))
    assert r["status"] == cfg.STATUS_BROKEN
    assert r["broken_reason"] == cfg.BROKEN_CONTAINER_DELETED
    # тег есть, сеть молчит -> загрузка не зафиксирована
    r = _run_worker(FakePage([], GTM_HTML))
    assert r["broken_reason"] == cfg.BROKEN_NOT_LOADED
    # у WORKING причины нет
    r = _run_worker(FakePage([COLLECT], GTM_HTML))
    assert r["broken_reason"] is None


def test_targeted_wait():
    """Гонка gudok: тег в вёрстке есть, запрос к gtm.js «опаздывает».
    Прицельное дожидание ловит его -> точная причина вместо not_loaded."""
    page = FakePage([], GTM_HTML, late_responses=[GTM404])
    r = _run_worker(page)
    assert r["status"] == cfg.STATUS_BROKEN
    assert r["broken_reason"] == cfg.BROKEN_CONTAINER_DELETED  # поймали 404
    assert any("googletagmanager.com" in q["url"] for q in r["ga_requests"])

    # сильного маркера в вёрстке нет -> дожидание не запускается,
    # «опоздавшие» ответы никто не ждёт
    page = FakePage([], WEAK_HTML, late_responses=[GTM404])
    r = _run_worker(page)
    assert r["status"] == cfg.STATUS_NONE
    assert not r["ga_requests"]


def test_probe():
    """Зонд: страница не запрашивает контейнер — чекер добывает улику сам."""
    from modules.page_worker import _extract_probe_target

    # выбор цели: GTM-контейнер приоритетнее потока G-, «сырые» маркеры чистятся
    target = _extract_probe_target([{"marker": "id=G-XYZ98765"},
                                    {"marker": "GTM-AB12CD"}])
    assert target == ("GTM-AB12CD",
                      "https://www.googletagmanager.com/gtm.js?id=GTM-AB12CD")
    cid, purl = _extract_probe_target([{"marker": "id=G-XYZ98765"}])
    assert cid == "G-XYZ98765" and purl.endswith("gtag/js?id=G-XYZ98765")

    # gudok при полном молчании сети: зонд 404 -> контейнер удалён
    page = FakePage([], GTM_HTML, probe_status=404)
    r = _run_worker(page)
    assert r["status"] == cfg.STATUS_BROKEN
    assert r["broken_reason"] == cfg.BROKEN_CONTAINER_DELETED
    assert r["probe"]["status"] == 404
    assert page.probe_calls == [r["probe"]["url"]]  # ровно один GET

    # контейнер жив (200), но страница его не запрашивает
    page = FakePage([], GTM_HTML, probe_status=200)
    r = _run_worker(page)
    assert r["broken_reason"] == cfg.BROKEN_ALIVE_NOT_LOADED
    assert r["probe"]["status"] == 200

    # страница сама сходила за контейнером -> зонд не запускается
    page = FakePage([GTM404], GTM_HTML, probe_status=200)
    r = _run_worker(page)
    assert r["probe"] is None and not page.probe_calls
    assert r["broken_reason"] == cfg.BROKEN_CONTAINER_DELETED  # улика страницы

    # хиты летят (WORKING) -> зонд не нужен даже без запроса контейнера
    page = FakePage([COLLECT], GTM_HTML, probe_status=404)
    r = _run_worker(page)
    assert r["status"] == cfg.STATUS_WORKING
    assert r["probe"] is None and not page.probe_calls

    # без ID в вёрстке зонд не применяется вовсе
    page = FakePage([], WEAK_HTML, probe_status=404)
    r = _run_worker(page)
    assert r["probe"] is None and not page.probe_calls

    # ошибка зонда не роняет прогон и не выдумывает улик
    page = FakePage([], GTM_HTML, probe_error="net down")
    r = _run_worker(page)
    assert r["status"] == cfg.STATUS_BROKEN
    assert r["broken_reason"] == cfg.BROKEN_NOT_LOADED
    assert r["probe"]["error"] == "net down"


def test_summary_report():
    """Краткий отчёт «для чата»: обнаружено/не обнаружено/ошибки по видам."""
    results = [
        _run_worker(FakePage([COLLECT], GTM_HTML)),           # хиты -> found
        _run_worker(FakePage([GTM404, ERRTEL], GTM_HTML)),    # gtm.js -> found
        _run_worker(FakePage([ERRTEL], CLEAN_HTML)),          # телеметрия -> off
        _run_worker(FakePage([], WEAK_HTML)),                 # слабые -> off
        _run_worker(FakePage([], CLEAN_HTML)),                # чисто -> off
        # страничного запроса нет, улику добыл только зонд ->
        # в мини-отчёте это «не обнаружено» (зонд — не network страницы)
        _run_worker(FakePage([], GTM_HTML, probe_status=404)),
        _run_worker(FakePage([], CLEAN_HTML, fail=True)),     # ERROR
    ]
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / "summary.txt"
        writer = ReportWriter(Path(tmp) / "full.txt")
        text = writer.write_summary(results, out)
        assert out.read_text(encoding="utf-8") == text  # файл = текст

    for fragment in ("=== ОТЧЁТ ПО НАЛИЧИЮ GOOGLE ANALYTICS ===",
                     "Всего URL в списке          : 7",
                     "Успешно открыто (всего)     : 6",
                     "  └─ Google обнаружено      : 2",
                     "  └─ Google не обнаружено   : 4",
                     "Ошибок (не удалось открыть) : 1",
                     "--- Google обнаружен ---",
                     "Статус : Google Found",
                     "--- Ошибка при открытии ---",
                     # ошибки сгруппированы по причине, а не свалены списком
                     "домен не резолвится",
                     "КОНЕЦ ОТЧЁТА"):
        assert fragment in text, f"в кратком отчёте нет: {fragment}"

    # расхождение с расширенным отчётом объясняется прямо в шапке:
    # сайт с зондовой уликой — это «тег есть, но не срабатывает»
    assert "тег GA лежит в вёрстке, но не срабатывает : 1" in text

    # статусы и причины в кратком отчёте не подсвечиваются
    assert "BROKEN" not in text and "WORKING" not in text
    assert "SUSPICIOUS" not in text and "контекст:" not in text
    assert "зонд" not in text
    assert "--- Google не обнаружен" not in text
    assert text.count("Google Found") == 2
    # компактная структура: без пустых строк между записями
    assert "\n\n" not in text


def test_excel_reader():
    """Мультилистовое чтение: гэпы <25 проходим, 25+ — переход на лист."""
    try:
        import openpyxl
    except ImportError:
        print("SKIP test_excel_reader: нет openpyxl")
        return
    from modules.excel_reader import read_urls_from_excel

    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Лист1"
    ws1.append(["url"])
    ws1.append(["https://a.ru"])
    for _ in range(3):                # маленький разрыв — проходим насквозь
        ws1.append([None])
    ws1.append(["https://b.ru"])
    for _ in range(25):               # 25 пустых подряд — конец листа
        ws1.append([None])
    ws1.append(["https://hidden.ru"])  # за разрывом — НЕ читается

    ws2 = wb.create_sheet("Лист2")    # второй лист тоже читается
    ws2.append(["url"])
    ws2.append(["https://c.ru"])
    ws2.append(["https://a.ru"])      # дубль с первого листа — отсеется
    # «грязные» ячейки — реальные случаи из боевого списка:
    ws2.append(["  https://spaced.ru  "])                 # лишние пробелы
    ws2.append(["https://travel.example.com / tg.example.com"])       # два URL в ячейке
    ws2.append(["one.ru, two.ru; three.ru"])              # список в ячейке
    ws2.append(["https://other pay.example.com"])            # мусорное слово + URL
    ws2.append(["https://xn--80aswg.xn--p1ai"])           # punycode-домен
    ws2.append(["https://query.ru?utm=1"])                # query без пути
    ws2.append(["lk b2b"])                                # домена нет вовсе
    # дубли в другой записи: без схемы / в другом регистре / список
    # «запятая или точка с запятой С ПРОБЕЛОМ» — ничего нового не дадут
    ws2.append(["a.ru; ONE.RU, https://TWO.ru"])

    ws3 = wb.create_sheet("Служебный")  # листа без колонки url — пропуск
    ws3.append(["комментарий"])
    ws3.append(["не трогать"])

    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "input.xlsx"
        wb.save(path)

        urls = list(read_urls_from_excel(path))
        assert urls == [
            "https://a.ru", "https://b.ru",           # Лист1 до разрыва
            "https://c.ru",                           # Лист2 (a.ru — дубль)
            "https://spaced.ru",                      # пробелы обрезаны
            "https://travel.example.com", "tg.example.com",       # ячейка разделена
            "one.ru", "two.ru", "three.ru",           # список разделён
            "pay.example.com",                           # мусорное слово убрано
            "https://xn--80aswg.xn--p1ai",            # punycode не обрезан
            "https://query.ru?utm=1",                 # query сохранён
            "lkb2b",                                  # без домена: до браузера
        ], urls                                       # не дойдёт (invalid_url)

        # старое поведение: конкретный лист по имени
        only2 = list(read_urls_from_excel(path, sheet_name="Лист2"))
        assert only2[:2] == ["https://c.ru", "https://a.ru"]

        # нет колонки ни на одном листе -> ошибка
        try:
            list(read_urls_from_excel(path, column_name="нет_такой"))
            assert False, "ожидали ValueError"
        except ValueError:
            pass


def test_aggregator():
    """Слияние прогонов: лучший результат на URL, без дублей, порядок."""
    from modules.aggregator import merge_runs, has_network_evidence

    def res(url, ga=None, collect=None, error=None):
        return {"url": url, "status": cfg.STATUS_ERROR if error else "NONE",
                "ga_requests": ga or [], "collect_requests": collect or [],
                "error": error}

    gtm = {"url": "https://www.googletagmanager.com/gtm.js?id=GTM-AB12CD",
           "status": 404, "kind": "other_google"}
    hit = {"url": "https://www.google-analytics.com/g/collect?v=2",
           "status": 204, "kind": "collect"}

    # критерий сетевой улики — общий с мини-отчётом
    assert has_network_evidence(res("x", ga=[gtm]))
    assert has_network_evidence(res("x", collect=[hit]))
    assert not has_network_evidence(res("x"))
    assert not has_network_evidence(
        res("x", ga=[{"url": "https://www.google-analytics.com/analytics.js"}]))

    run1 = [res("https://a.ru", error="net::ERR_TIMED_OUT"),  # ошибка
            res("https://b.ru"),                              # пусто
            res("https://c.ru", ga=[gtm]),                    # контейнер
            res("https://e.ru", error="net::ERR_TIMED_OUT")]  # ошибка
    run2 = [res("https://a.ru", collect=[hit]),               # хиты!
            res("https://b.ru", ga=[gtm]),                    # контейнер!
            res("https://c.ru", error="net::ERR_TIMED_OUT"),  # ошибка
            res("https://d.ru"),                              # новый URL
            res("https://e.ru", error="net::ERR_REFUSED")]    # снова ошибка

    merged = merge_runs([run1, run2])
    by_url = {r["url"]: r for r in merged}

    # без дублей, порядок первого прогона + новые в конце
    assert [r["url"] for r in merged] == [
        "https://a.ru", "https://b.ru", "https://c.ru",
        "https://e.ru", "https://d.ru"]
    # ошибка проигрывает любому открытию: a — хиты, c — контейнер из run1
    assert by_url["https://a.ru"]["collect_requests"]
    assert by_url["https://b.ru"]["ga_requests"]
    assert by_url["https://c.ru"]["ga_requests"] and \
        not by_url["https://c.ru"]["error"]
    # ошибка во ВСЕХ прогонах — остаётся ошибкой (берётся первая)
    assert by_url["https://e.ru"]["error"] == "net::ERR_TIMED_OUT"
    # счётчик нестабильности: e.ru не открылся ни разу из двух прогонов
    assert by_url["https://e.ru"]["runs_failed"] == 2
    assert by_url["https://a.ru"]["runs_failed"] == 1

    # общий мини-отчёт по слитым результатам: 3 found, 1 не обнаружено, 1 err
    with tempfile.TemporaryDirectory() as tmp:
        writer = ReportWriter(Path(tmp) / "full.txt")
        text = writer.build_summary(merged)
    for fragment in ("Всего URL в списке          : 5",
                     "  └─ Google обнаружено      : 3",
                     "  └─ Google не обнаружено   : 1",
                     "Ошибок (не удалось открыть) : 1"):
        assert fragment in text, f"в общем мини-отчёте нет: {fragment}"


def test_aggregator_markup_rank():
    """Маркеры в вёрстке — отдельный ярус силы улики.

    Раньше BROKEN и NONE делили один ранг, а замена шла только при
    строгом превосходстве: классификацию фиксировал ПЕРВЫЙ успешный
    прогон. Если в нём не прочитался HTML, сайт навсегда оставался NONE,
    сколько бы прогонов ни делали — то есть 19 из 20 прогонов боевого
    запуска не могли исправить ровно тот случай, ради которого их
    и делали.
    """
    from modules.aggregator import merge_runs, has_strong_markers

    strong = [{"marker": "GTM-AB12CD", "strength": "strong", "snippet": "x"}]

    def res(url, markers=None, status="NONE", http_error=False):
        return {"url": url, "status": status, "ga_requests": [],
                "collect_requests": [], "html_markers": markers or [],
                "http_error": http_error, "error": None}

    assert has_strong_markers(res("x", strong))
    assert not has_strong_markers(res("x"))

    # прогон 1 не увидел вёрстку, прогон 2 увидел -> побеждает прогон 2
    run1 = [res("https://a.ru")]
    run2 = [res("https://a.ru", strong, status=cfg.STATUS_BROKEN)]
    merged = merge_runs([run1, run2])
    assert merged[0]["status"] == cfg.STATUS_BROKEN, \
        "поздний прогон с маркерами должен вытеснять пустой результат"

    # обратный порядок: лучший результат уже первый, замены нет
    merged = merge_runs([run2, run1])
    assert merged[0]["status"] == cfg.STATUS_BROKEN

    # при равных уликах чистый ответ сервера лучше кода ошибки
    run1 = [res("https://b.ru", strong, cfg.STATUS_BROKEN, http_error=True)]
    run2 = [res("https://b.ru", strong, cfg.STATUS_BROKEN)]
    assert not merge_runs([run1, run2])[0]["http_error"]


def test_notifier():
    """Mattermost-вебхук: успех, ошибка сети (без падения), пустой URL."""
    import urllib.error
    from modules import notifier

    class FakeResponse:
        status = 200
        def __enter__(self):
            return self
        def __exit__(self, *args):
            return False

    calls = {}

    def fake_urlopen_ok(request, timeout=None):
        calls["url"] = request.full_url
        calls["body"] = request.data.decode("utf-8")
        return FakeResponse()

    def fake_urlopen_fail(request, timeout=None):
        raise urllib.error.URLError("connection refused")

    original = notifier.urllib.request.urlopen
    try:
        notifier.urllib.request.urlopen = fake_urlopen_ok
        assert notifier.send_to_mattermost("https://mm.local/hooks/abc",
                                           "текст отчёта") is True
        assert calls["url"] == "https://mm.local/hooks/abc"
        assert "текст отчёта" in calls["body"]        # ensure_ascii=False

        notifier.urllib.request.urlopen = fake_urlopen_fail
        assert notifier.send_to_mattermost("https://mm.local/hooks/abc",
                                           "x") is False  # не упало
    finally:
        notifier.urllib.request.urlopen = original

    assert notifier.send_to_mattermost("", "x") is False  # выключено


def test_browser_manager_semaphore():
    """Семафор ограничивает число ОДНОВРЕМЕННО ОТКРЫТЫХ страниц."""
    from modules.browser_manager import BrowserManager

    limit = 2
    bm = BrowserManager(max_concurrent_pages=limit)
    state = {"active": 0, "peak": 0}

    async def fake_new_page():
        state["active"] += 1
        state["peak"] = max(state["peak"], state["active"])
        return object()

    async def fake_close_page(page):
        state["active"] -= 1

    bm.new_page = fake_new_page       # подменяем реальный браузер
    bm.close_page = fake_close_page

    async def use_page():
        async with bm.page_context():
            await asyncio.sleep(0.01)  # «работаем» со страницей

    async def run_many():
        await asyncio.gather(*(use_page() for _ in range(10)))

    asyncio.run(run_many())
    assert state["active"] == 0                     # всё закрыто
    assert state["peak"] <= limit, state["peak"]    # лимит не превышен
    assert state["peak"] == limit                   # и при этом использован


def test_main_error_fallback():
    """Аварийная ветка оркестратора и детект мёртвого браузера.

    Если Chromium умирает посреди прогона, каждая страница падает
    на new_page. Тогда важно (а) записать НОРМАЛИЗОВАННЫЙ адрес —
    иначе слияние прогонов сочтёт «lk.example.com» и «https://lk.example.com»
    разными сайтами и раздует отчёт дублями; (б) не отрабатывать
    оставшиеся прогоны на мёртвом браузере.
    """
    import main as m
    from modules.browser_manager import BrowserManager

    class DeadBM:
        @asynccontextmanager
        async def page_context(self):
            raise RuntimeError("Target page, context or browser has been closed")
            yield  # pragma: no cover

    results = asyncio.run(m.process_urls(["lk.example.com", "b.example.com"], DeadBM()))
    assert [r["url"] for r in results] == ["https://lk.example.com", "https://b.example.com"]
    assert all(r["status"] == cfg.STATUS_ERROR for r in results)
    # порядок совпадает с входным файлом, а не с порядком завершения задач
    merged = __import__("modules.aggregator", fromlist=["merge_runs"]).merge_runs(
        [results, results])
    assert len(merged) == 2, "нормализация адреса не спасла от дублей"

    bm = BrowserManager()
    assert not bm.is_alive()                       # браузер не запускался
    bm._browser = types.SimpleNamespace(is_connected=lambda: True)
    assert bm.is_alive()
    bm._browser = types.SimpleNamespace(is_connected=lambda: False)
    assert not bm.is_alive()


def test_browser_launch_args():
    """Устаревший TLS разрешён — иначе сайты на TLS 1.0/1.1 не открыть."""
    from modules.browser_manager import BrowserManager

    args = BrowserManager(allow_legacy_tls=True, headless=True)._launch_args()
    assert "--ssl-version-min=tls1" in args
    assert "--ignore-certificate-errors" in args
    assert "--headless=new" in args

    args = BrowserManager(allow_legacy_tls=False, headless=False)._launch_args()
    assert "--ssl-version-min=tls1" not in args
    assert "--headless=new" not in args

    # точка расширения под разделение на контуры
    args = BrowserManager(extra_launch_args=["--host-resolver-rules=X"])._launch_args()
    assert "--host-resolver-rules=X" in args


def main():
    tests = [
        test_config_patterns,
        test_classifier,
        test_failures_taxonomy,
        test_url_variants,
        test_network_listener,
        test_network_listener_broken_response,
        test_find_html_markers,
        test_page_worker,
        test_no_retry_on_permanent_failure,
        test_alternate_url_ladder,
        test_page_reset_between_variants,
        test_http_error_pages_are_analyzed,
        test_report_writer,
        test_container_grouping,
        test_universal_analytics_is_dead,
        test_ua_in_reports,
        test_mattermost_webhook_source,
        test_broken_reason,
        test_targeted_wait,
        test_probe,
        test_summary_report,
        test_excel_reader,
        test_aggregator,
        test_aggregator_markup_rank,
        test_notifier,
        test_browser_manager_semaphore,
        test_browser_launch_args,
        test_main_error_fallback,
    ]
    failed = 0
    for test in tests:
        try:
            test()
        except AssertionError as e:
            failed += 1
            print(f"FAIL {test.__name__}: {e}")
        else:
            print(f"OK  {test.__name__}")
    if failed:
        print(f"\nПровалено тестов: {failed} из {len(tests)}.")
        sys.exit(1)
    print(f"\nВсе {len(tests)} тестов пройдены.")


if __name__ == "__main__":
    main()
