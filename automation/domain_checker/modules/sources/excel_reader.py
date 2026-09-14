# Читалка списка адресов из Excel. Ведёт начало от analytics/google_page_checker,
# но переписана на голый openpyxl: pandas тянулся сюда ради двух вызовов, а внутри
# всё равно работал openpyxl. Для скрипта, который ставится на хост одной командой,
# это была самая тяжёлая зависимость в requirements.
from __future__ import annotations

import logging
import re
from pathlib import Path
from typing import Generator, List

from openpyxl import load_workbook

logger = logging.getLogger(__name__)


def _is_empty(value) -> bool:
    return value is None or not str(value).strip()

# Доменная зона — либо буквенная (ru, com, рф), либо punycode (xn--p1ai).
# Без варианта xn-- живой адрес https://xn--80aswg.xn--p1ai обрезался до
# "https://xn--80aswg.xn" и стабильно уходил в ERROR как несуществующий.
# Хвост начинается с /, ? или #, иначе у адресов без пути терялся query.
_URL_TOKEN_RE = re.compile(
    r"(?:https?://)?(?:[\w-]+\.)+(?:xn--[a-z0-9-]+|[A-Za-zА-Яа-я]{2,})"
    r"(?::\d+)?(?:[/?#]\S*)?",
    re.IGNORECASE,
)


def _extract_urls(raw) -> List[str]:
    """
    Если ни одного URL-токена не нашлось (в ячейке нет ни одного слова
    с доменом), возвращаем содержимое одной строкой БЕЗ пробелов —
    такая запись честно доедет до отсева и попадёт в skipped,
    а не потеряется молча.
    """
    text = str(raw).strip()
    found = [m.group(0).rstrip(",;.") for m in _URL_TOKEN_RE.finditer(text)]
    if found:
        return found
    collapsed = "".join(text.split())
    return [collapsed] if collapsed else []


def _pick_sheets(workbook, sheet_name) -> list:
    if sheet_name is None:
        return list(workbook.worksheets)
    if isinstance(sheet_name, int):
        return [workbook.worksheets[sheet_name]]
    return [workbook[sheet_name]]


def _header_map(sheet) -> dict:
    """Заголовок это первая строка листа: имя колонки -> её номер."""
    for row in sheet.iter_rows(min_row=1, max_row=1, values_only=True):
        return {str(v).strip(): i for i, v in enumerate(row)
                if not _is_empty(v)}
    return {}


def read_urls_from_excel(
    file_path: str | Path,
    column_name: str = "url",
    sheet_name: int | str | None = None,
    empty_row_break: int = 25,
    extra_column: str = None,
) -> Generator:
    """Читает адреса из Excel, лист за листом.

    Без extra_column отдаёт строки, с extra_column — пары (адрес, значение
    соседней колонки той же строки): так к адресу приезжает метка контура.
    Колонки нет на листе — значение приходит пустым.

    Дедупликации здесь нет намеренно: она живёт в targets.py, и место,
    где домен может пропасть, должно быть одно.
    """
    file_path = Path(file_path)
    if not file_path.exists():
        raise FileNotFoundError("Excel-файл не найден: {}".format(file_path))

    logger.info(
        "Чтение URL из файла: %s (листы: %s, колонка: %s)",
        file_path, "все" if sheet_name is None else sheet_name, column_name,
    )

    # read_only держит лист потоком, data_only отдаёт посчитанные значения
    # вместо текста формул.
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    try:
        sheets = _pick_sheets(workbook, sheet_name)
        headers = {sheet.title: _header_map(sheet) for sheet in sheets}

        # Колонка должна найтись хотя бы на одном листе — иначе это не тот файл.
        if not any(column_name in h for h in headers.values()):
            available = ", ".join(
                "{}: [{}]".format(title, ", ".join(h))
                for title, h in headers.items())
            raise ValueError(
                "Колонка '{}' не найдена ни на одном листе. "
                "Доступные колонки: {}".format(column_name, available))

        total = 0
        for sheet in sheets:
            header = headers[sheet.title]
            if column_name not in header:
                logger.warning("Лист '%s': колонки '%s' нет — лист пропущен",
                               sheet.title, column_name)
                continue

            url_col = header[column_name]
            extra_col = header.get(extra_column) if extra_column else None
            sheet_count = 0
            empty_streak = 0
            stopped_by_gap = False

            for row in sheet.iter_rows(min_row=2, values_only=True):
                value = row[url_col] if url_col < len(row) else None
                if _is_empty(value):
                    empty_streak += 1
                    if empty_streak >= empty_row_break:
                        # Длинный разрыв — данные листа закончились,
                        # переходим к следующему листу.
                        stopped_by_gap = True
                        break
                    continue
                empty_streak = 0

                extra = None
                if extra_col is not None and extra_col < len(row):
                    extra = row[extra_col]

                # Ячейка может содержать несколько адресов и лишние пробелы —
                # разбираем её на отдельные.
                for url in _extract_urls(value):
                    sheet_count += 1
                    total += 1
                    if extra_column:
                        yield url, ("" if _is_empty(extra)
                                    else str(extra).strip())
                    else:
                        yield url

            logger.info(
                "Лист '%s': прочитано URL: %d (%s)", sheet.title, sheet_count,
                "остановлено разрывом {}+ пустых строк".format(empty_row_break)
                if stopped_by_gap else "дочитан до конца",
            )

        logger.info("Завершено чтение URL со всех листов. Всего: %d", total)
    finally:
        # read_only оставляет открытый файловый дескриптор.
        workbook.close()
